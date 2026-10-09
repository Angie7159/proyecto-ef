from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
import json
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


def _format_date(value: Any) -> str:
    if value is None or value == "":
        return "Sin fecha"
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    try:
        return date.fromisoformat(str(value)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return str(value)


def _set_word_font(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        document.styles[style_name].font.name = "Calibri"


def _save_document(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _agregar_titulo_seccion(doc: Document, texto: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(texto)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return paragraph


def _agregar_parrafo(doc: Document, texto: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(texto)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return paragraph


def generar_word_resumen(reunion: dict[str, Any], acuerdos: list[dict[str, Any]] | None = None) -> bytes:
    """Crea un resumen ejecutivo Word estructurado en ocho secciones."""
    acuerdos = acuerdos or []
    document = Document()
    for section in document.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
    _set_word_font(document)

    title = reunion.get("titulo_final") or "Reunión sin título"
    meeting_date = reunion.get("fecha")
    if isinstance(meeting_date, (date, datetime)):
        date_display = meeting_date.strftime("%d/%m/%Y")
    elif meeting_date:
        try:
            date_display = datetime.fromisoformat(str(meeting_date).replace("Z", "+00:00")).strftime("%d/%m/%Y")
        except ValueError:
            date_display = str(meeting_date)
    else:
        date_display = ""

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(3)
    title_run = title_paragraph.add_run(f"{title.upper()} — {date_display}".strip())
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    divider = document.add_paragraph()
    divider.paragraph_format.space_after = Pt(12)
    paragraph_properties = divider._p.get_or_add_pPr()
    paragraph_borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), "1E3A8A")
    paragraph_borders.append(bottom)
    paragraph_properties.append(paragraph_borders)

    sections = [
        ("Resumen Ejecutivo", reunion.get("resumen_ejecutivo") or "Resumen no disponible."),
        ("Problema", reunion.get("seccion_problema") or "No se identificó un problema específico en la reunión."),
        ("Solución", reunion.get("seccion_solucion") or "No se presentó una solución específica en la reunión."),
        ("Mercado", reunion.get("seccion_mercado") or "No se discutió contexto de mercado en esta reunión."),
        ("Finanzas", reunion.get("seccion_finanzas") or "No se presentaron cifras económicas en esta reunión."),
        ("Conclusión", reunion.get("seccion_conclusion") or "No se registró una conclusión específica en la reunión."),
    ]
    for heading, content in sections:
        _agregar_titulo_seccion(document, heading)
        _agregar_parrafo(document, str(content))

    raw_steps = reunion.get("seccion_proximos_pasos") or []
    if isinstance(raw_steps, str):
        try:
            steps = json.loads(raw_steps)
        except json.JSONDecodeError:
            steps = [raw_steps] if raw_steps.strip() else []
    else:
        steps = raw_steps
    if not isinstance(steps, list):
        steps = [str(steps)] if steps else []
    _agregar_titulo_seccion(document, "Próximos pasos")
    if steps:
        for index, step in enumerate(steps, start=1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.keep_together = True
            number = paragraph.add_run(f"({index}) ")
            number.font.name = "Calibri"
            number.font.size = Pt(11)
            number.font.bold = True
            text = paragraph.add_run(str(step))
            text.font.name = "Calibri"
            text.font.size = Pt(11)
    else:
        _agregar_parrafo(document, "No se identificaron próximos pasos específicos.")

    participants = reunion.get("participantes_lista") or []
    if participants:
        document.add_page_break()
        _agregar_titulo_seccion(document, "Participantes")
        for participant in participants:
            if isinstance(participant, dict):
                name = participant.get("nombre") or participant.get("name") or ""
                role = participant.get("cargo") or participant.get("rol") or participant.get("role") or ""
            else:
                name, role = str(participant), ""
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            name_run = paragraph.add_run(f"• {name}")
            name_run.font.name = "Calibri"
            name_run.font.size = Pt(10)
            name_run.font.bold = True
            if role:
                role_run = paragraph.add_run(f" — {role}")
                role_run.font.name = "Calibri"
                role_run.font.size = Pt(10)
                role_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    return _save_document(document)

def _set_cell_background(cell, color_hex: str) -> None:
    """Aplica un relleno sÃ³lido a una celda de tabla Word."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color_hex)


def _set_cell_borders(cell, color_hex: str = "CBD5E1", size: int = 6) -> None:
    """Aplica bordes consistentes a los cuatro lados de una celda."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{side}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{side}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:color"), color_hex)


def _short_text(value: Any, limit: int = 135) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: max(0, limit - 3)].rstrip() + "..."


def _map_item_parts(item: Any, kind: str, index: int) -> tuple[str, str]:
    if isinstance(item, dict):
        if kind == "tema":
            title = item.get("titulo") or item.get("tema") or item.get("nombre") or f"Tema {index}"
            description = item.get("descripcion") or item.get("detalle") or ""
        elif kind == "debate":
            title = item.get("titulo") or item.get("problema") or f"ObservaciÃ³n {index}"
            description = item.get("descripcion") or item.get("detalle") or ""
        elif kind == "acuerdo":
            title = item.get("titulo") or f"Acuerdo {index}"
            description = item.get("descripcion") or item.get("detalle") or ""
        else:
            title = item.get("titulo") or item.get("descripcion") or f"Tarea {index}"
            description = item.get("detalle") or ""
            owner = item.get("responsable") or "Sin responsable"
            deadline = _format_date(item.get("fecha_limite")) if item.get("fecha_limite") else "Sin plazo"
            description = " Â· ".join(part for part in (description, f"Responsable: {owner}", f"Plazo: {deadline}") if part)
        return _short_text(title, 70), _short_text(description, 145)
    if kind in {"tema", "debate"}:
        label = "Tema" if kind == "tema" else "ObservaciÃ³n"
        return f"{label} {index}", _short_text(item, 145)
    return _short_text(item, 135), ""


def _style_card_paragraph(paragraph, accent: str, background: str) -> None:
    paragraph.paragraph_format.left_indent = Pt(4)
    paragraph.paragraph_format.right_indent = Pt(3)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), background)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), accent)
        borders.append(border)
    p_pr.append(borders)


def _add_map_card(cell, title: str, description: str, accent: str, background: str) -> None:
    title_paragraph = cell.add_paragraph()
    _style_card_paragraph(title_paragraph, accent, background)
    heading = title_paragraph.add_run(f"â€¢ {_short_text(title, 70) or 'Elemento'}")
    heading.font.name = "Calibri"
    heading.font.size = Pt(9)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(accent)
    if description:
        description_paragraph = cell.add_paragraph()
        _style_card_paragraph(description_paragraph, accent, background)
        description_run = description_paragraph.add_run(_short_text(description, 145))
        description_run.font.name = "Calibri"
        description_run.font.size = Pt(8)
        description_run.font.italic = True
        description_run.font.color.rgb = RGBColor.from_string("374151")


def generar_word_mapa_conceptual(
    reunion: dict[str, Any], acuerdos: list[dict[str, Any]] | None = None
) -> bytes:
    """Crea un mapa conceptual editable con tablas nativas de Word."""
    acuerdos = acuerdos or []
    document = Document()
    _set_word_font(document)

    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)

    committee = reunion.get("comites") or {}
    if isinstance(committee, list):
        committee = committee[0] if committee else {}
    committee_name = reunion.get("comite_nombre") or (
        committee.get("nombre") if isinstance(committee, dict) else None
    ) or "REUNIÃ“N GENERAL"
    date_label = _format_date(reunion.get("fecha"))
    meeting_title = reunion.get("titulo_final") or reunion.get("titulo") or "ReuniÃ³n"

    title_table = document.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_table.autofit = False
    title_table.columns[0].width = Cm(18)
    title_cell = title_table.cell(0, 0)
    title_cell.width = Cm(18)
    title_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_cell_background(title_cell, "1E3A8A")
    _set_cell_borders(title_cell, "1E3A8A", size=10)
    first = title_cell.paragraphs[0]
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.paragraph_format.space_after = Pt(3)
    run = first.add_run(f"{committee_name.upper()} â€” {date_label}")
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    subtitle = title_cell.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(2)
    title_run = subtitle.add_run(_short_text(meeting_title, 120))
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(10)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(224, 231, 255)

    summary = _short_text(reunion.get("resumen_ejecutivo"), 250)
    if summary:
        objective = title_cell.add_paragraph()
        objective.alignment = WD_ALIGN_PARAGRAPH.CENTER
        objective.paragraph_format.space_after = Pt(1)
        objective_run = objective.add_run(f"Objetivo: {summary}")
        objective_run.font.name = "Calibri"
        objective_run.font.size = Pt(8)
        objective_run.font.italic = True
        objective_run.font.color.rgb = RGBColor(224, 231, 255)

    document.add_paragraph().paragraph_format.space_after = Pt(2)

    columns = [
        {"title": "1. TEMAS", "subtitle": "Puntos principales presentados", "items": reunion.get("puntos_clave") or [], "kind": "tema", "accent": "1E3A8A", "fill": "DBEAFE", "card": "EFF6FF"},
        {"title": "2. DEBATE", "subtitle": "Observaciones y problemas", "items": reunion.get("problemas_identificados") or [], "kind": "debate", "accent": "B45309", "fill": "FEF3C7", "card": "FFFBEB"},
        {"title": "3. ACUERDOS", "subtitle": "Decisiones finales tomadas", "items": acuerdos, "kind": "acuerdo", "accent": "047857", "fill": "D1FAE5", "card": "ECFDF5"},
        {"title": "4. TAREAS", "subtitle": "Acciones, responsables y plazos", "items": acuerdos, "kind": "tarea", "accent": "7C3AED", "fill": "EDE9FE", "card": "F5F3FF"},
    ]
    columns_table = document.add_table(rows=1, cols=4)
    columns_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    columns_table.autofit = False
    for table_column in columns_table.columns:
        table_column.width = Cm(4.4)
    for index, config in enumerate(columns):
        column_cell = columns_table.cell(0, index)
        column_cell.width = Cm(4.4)
        column_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        _set_cell_borders(column_cell, config["accent"], size=8)
        _set_cell_background(column_cell, "FFFFFF")
        header = column_cell.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.paragraph_format.space_after = Pt(2)
        heading = header.add_run(config["title"])
        heading.font.name = "Calibri"
        heading.font.size = Pt(9)
        heading.font.bold = True
        heading.font.color.rgb = RGBColor.from_string(config["accent"])
        sub = column_cell.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.space_after = Pt(5)
        sub_run = sub.add_run(config["subtitle"])
        sub_run.font.name = "Calibri"
        sub_run.font.size = Pt(6)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 116, 139)

        items = config["items"]
        if not items:
            _add_map_card(column_cell, "Sin elementos registrados", "", config["accent"], config["fill"])
            continue
        for item_index, item in enumerate(items, start=1):
            title, description = _map_item_parts(item, config["kind"], item_index)
            _add_map_card(column_cell, title, description, config["accent"], config["card"])

    document.add_paragraph().paragraph_format.space_after = Pt(1)
    instructions = document.add_paragraph()
    instructions.alignment = WD_ALIGN_PARAGRAPH.CENTER
    instructions.paragraph_format.space_before = Pt(4)
    instruction_run = instructions.add_run(
        "Para convertir el contenido a SmartArt: copia los tÃ­tulos y elementos de una columna, "
        "usa Inicio â†’ Convertir a SmartArt y elige un diseÃ±o de lista o jerarquÃ­a."
    )
    instruction_run.font.name = "Calibri"
    instruction_run.font.size = Pt(7)
    instruction_run.font.italic = True
    instruction_run.font.color.rgb = RGBColor(100, 116, 139)
    return _save_document(document)

def generar_word_mapa_mental(reunion: dict[str, Any]) -> bytes:
    """Compatibilidad con llamadas existentes; genera un mapa conceptual."""
    return generar_word_mapa_conceptual(reunion, [])

