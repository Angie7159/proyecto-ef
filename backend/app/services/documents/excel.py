"""Generación del libro consolidado de seguimiento por comité."""

from datetime import date, datetime
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1E3A8A"
GRAY = "374151"
PALE_GRAY = "F3F4F6"
PALE_RED = "FEE2E2"
PALE_BLUE = "DBEAFE"
PALE_YELLOW = "FEF3C7"
PALE_GREEN = "D1FAE5"
OVERDUE_RED = "FCA5A5"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADERS = [
    "N°", "Área / Responsable", "Tipo", "Tema / Actividad", "Descripción / Hallazgo",
    "Acción Comprometida", "Responsable Acción", "Fecha Límite", "Estado",
    "Actualización Semanal", "Semana Origen",
]


def _date(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    try:
        return date.fromisoformat(str(value)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return str(value)


def _short(value: Any, limit: int = 240) -> str:
    text = "" if value is None else str(value)
    return text if len(text) <= limit else text[: limit - 3].rstrip() + "..."


def _safe_sheet_name(value: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", value).strip(" '\u2019") or "Reunión"
    base = base[:31]
    name, suffix = base, 2
    while name.casefold() in used:
        ending = f" ({suffix})"
        name = base[: 31 - len(ending)] + ending
        suffix += 1
    used.add(name.casefold())
    return name


def _style_header(ws, row: int, start: int, end: int, color: str = NAVY) -> None:
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=start, max_col=end):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
    ws.row_dimensions[row].height = 32


def _style_body(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _meeting_rows(meeting: dict[str, Any], agreements: list[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    title = meeting.get("titulo_final") or meeting.get("titulo") or "Reunión"
    origin = _date(meeting.get("fecha"))
    for item in meeting.get("puntos_clave") or []:
        text = item.get("descripcion") or item.get("titulo") or str(item) if isinstance(item, dict) else str(item)
        rows.append([len(rows) + 1, "", "Informativo", _short(text, 150), _short(text), "", "", "", "", "", origin])
    for item in meeting.get("problemas_identificados") or []:
        if isinstance(item, dict):
            topic = item.get("titulo") or item.get("tema") or "Problema identificado"
            detail = item.get("descripcion") or item.get("hallazgo") or ""
            area = item.get("categoria") or ""
        else:
            topic, detail, area = "Problema identificado", str(item), ""
        rows.append([len(rows) + 1, area, "Problema", _short(topic, 150), _short(detail), "", "", "", "Pendiente", "", origin])
    for item in agreements:
        if not isinstance(item, dict):
            continue
        kind = item.get("tipo") or ("Acuerdo" if item.get("es_nuevo", True) else "Seguimiento")
        topic = item.get("tema") or item.get("titulo") or item.get("descripcion") or "Acuerdo"
        action = item.get("accion_comprometida") or item.get("descripcion") or ""
        status = str(item.get("estado") or "Pendiente").replace("_", " ").title()
        rows.append([
            len(rows) + 1, item.get("area") or item.get("categoria") or "",
            kind, _short(topic, 150), _short(item.get("hallazgo") or item.get("descripcion")),
            _short(action), item.get("responsable") or "Sin asignar", _date(item.get("fecha_limite")),
            status, _short(item.get("actualizacion_semanal") or ""),
            _date(item.get("semana_origen") or meeting.get("fecha")) or title,
        ])
    return rows


def _meeting_sheet(workbook: Workbook, meeting: dict[str, Any], committee: str, used: set[str]) -> list[list[Any]]:
    meeting_date = _date(meeting.get("fecha")) or "Sin fecha"
    try:
        parsed_date = date.fromisoformat(str(meeting.get("fecha"))[:10]).strftime("%d-%m-%Y")
    except (ValueError, TypeError):
        parsed_date = "Reunion"
    ws = workbook.create_sheet(_safe_sheet_name(parsed_date, used))
    title = meeting.get("titulo_final") or meeting.get("titulo") or "Reunión"
    ws.merge_cells("A1:K1")
    ws["A1"] = f"{committee} — {title}"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(bold=True, size=16, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 34
    participants = meeting.get("participantes_lista") or []
    participants_text = ", ".join(str(p.get("nombre") or p.get("name") or "") if isinstance(p, dict) else str(p) for p in participants)
    ws.merge_cells("A2:K2")
    ws["A2"] = f"Fecha: {meeting_date}    Participantes: {participants_text or meeting.get('participantes_count') or 'No especificados'}"
    ws["A2"].font = Font(color=GRAY, italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.append(HEADERS)
    _style_header(ws, 3, 1, len(HEADERS))
    rows = _meeting_rows(meeting, meeting.get("acuerdos") or [])
    for row in rows:
        ws.append(row)
        fill = {"Problema": PALE_RED, "Acuerdo": PALE_BLUE, "Seguimiento": PALE_YELLOW, "Informativo": PALE_GRAY}.get(str(row[2]), WHITE)
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=fill)
        status_color = {
            "pendiente": PALE_RED, "en proceso": PALE_YELLOW,
            "cumplido": PALE_GREEN, "vencido": OVERDUE_RED,
        }.get(str(row[8]).lower())
        if status_color:
            ws.cell(ws.max_row, 9).fill = PatternFill("solid", fgColor=status_color)
    if rows:
        _style_body(ws, 4, ws.max_row, 1, len(HEADERS))
        ws.auto_filter.ref = f"A3:K{ws.max_row}"
    ws.freeze_panes = "A4"
    return rows


def generar_excel_seguimiento(comite_nombre: str, reuniones: list[dict[str, Any]]) -> bytes:
    """Genera un libro consolidado con una pestaña por reunión y seguimiento."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    all_rows: list[list[Any]] = []
    committee = comite_nombre or "Comité"
    for meeting in reuniones:
        all_rows.extend(_meeting_sheet(workbook, meeting, committee, used))

    ws = workbook.create_sheet("SEGUIMIENTO")
    ws.append(["N°", "Reunión Origen", "Área", "Tipo", "Tema / Actividad", "Acción Comprometida", "Estado", "Responsable", "Fecha Límite", "Actualización Semanal", "Semana Origen"])
    _style_header(ws, 1, 1, 11)
    for meeting in reuniones:
        title = meeting.get("titulo_final") or meeting.get("titulo") or "Reunión"
        for row in _meeting_rows(meeting, meeting.get("acuerdos") or []):
            ws.append([row[0], title, row[1], row[2], row[3], row[5], row[8], row[6], row[7], row[9], row[10]])
            fill = {"Problema": PALE_RED, "Acuerdo": PALE_BLUE, "Seguimiento": PALE_YELLOW, "Informativo": PALE_GRAY}.get(str(row[2]), WHITE)
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
            status_color = {
                "pendiente": PALE_RED, "en proceso": PALE_YELLOW,
                "cumplido": PALE_GREEN, "vencido": OVERDUE_RED,
            }.get(str(row[8]).lower())
            if status_color:
                ws.cell(ws.max_row, 7).fill = PatternFill("solid", fgColor=status_color)
    if ws.max_row > 1:
        _style_body(ws, 2, ws.max_row, 1, 11)
        ws.auto_filter.ref = f"A1:K{ws.max_row}"
    ws.freeze_panes = "A2"
    summary_row = ws.max_row + 2
    ws.cell(summary_row, 1, "Indicadores")
    ws.cell(summary_row + 1, 1, "Pendientes / en proceso")
    ws.cell(summary_row + 1, 2, f'=COUNTIF(G2:G{max(ws.max_row, 2)},"Pendiente")+COUNTIF(G2:G{max(ws.max_row, 2)},"En Proceso")')
    ws.cell(summary_row + 2, 1, "Cumplidos")
    ws.cell(summary_row + 2, 2, f'=COUNTIF(G2:G{max(ws.max_row, 2)},"Cumplido")')
    ws.cell(summary_row + 3, 1, "Vencidos")
    ws.cell(summary_row + 3, 2, f'=COUNTIF(G2:G{max(ws.max_row, 2)},"Vencido")')
    ws.cell(summary_row, 1).font = Font(bold=True, color=WHITE)
    ws.cell(summary_row, 1).fill = PatternFill("solid", fgColor=NAVY)

    categories = sorted({str(row[1]) for row in all_rows if row[1]})
    if categories:
        kpi = workbook.create_sheet("KPIs x ÁREA")
        kpi.append(["Área", "Total de items", "Pendientes", "En proceso", "Cumplidos", "Vencidos"])
        _style_header(kpi, 1, 1, 6)
        for category in categories:
            kpi.append([category, f'=COUNTIF(SEGUIMIENTO!C:C,A{kpi.max_row + 1})', f'=COUNTIFS(SEGUIMIENTO!C:C,A{kpi.max_row + 1},SEGUIMIENTO!G:G,"Pendiente")', f'=COUNTIFS(SEGUIMIENTO!C:C,A{kpi.max_row + 1},SEGUIMIENTO!G:G,"En Proceso")', f'=COUNTIFS(SEGUIMIENTO!C:C,A{kpi.max_row + 1},SEGUIMIENTO!G:G,"Cumplido")', f'=COUNTIFS(SEGUIMIENTO!C:C,A{kpi.max_row + 1},SEGUIMIENTO!G:G,"Vencido")'])
        _style_body(kpi, 2, kpi.max_row, 1, 6)

    legend = workbook.create_sheet("LEYENDA")
    legend.append(["Tipo / Estado", "Color", "Descripción"])
    _style_header(legend, 1, 1, 3)
    for label, color, description in [
        ("Informativo", PALE_GRAY, "Punto clave tratado"), ("Problema", PALE_RED, "Hallazgo o dificultad"),
        ("Acuerdo", PALE_BLUE, "Decisión o compromiso"), ("Seguimiento", PALE_YELLOW, "Acción de seguimiento"),
        ("Pendiente", PALE_RED, "Aún no iniciado"), ("En proceso", PALE_YELLOW, "En ejecución"),
        ("Cumplido", PALE_GREEN, "Finalizado"), ("Vencido", OVERDUE_RED, "Fuera de plazo"),
    ]:
        legend.append([label, color, description])
        legend.cell(legend.max_row, 1).fill = PatternFill("solid", fgColor=color)
    _style_body(legend, 2, legend.max_row, 1, 3)

    widths = {"A": 10, "B": 24, "C": 18, "D": 27, "E": 42, "F": 42, "G": 23, "H": 18, "I": 16, "J": 35, "K": 18}
    for sheet in workbook.worksheets:
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        for row in range(1, sheet.max_row + 1):
            if row > 3:
                sheet.row_dimensions[row].height = 38
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
